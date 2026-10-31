import os
import io
import json
import csv
from typing import Dict, Any, List

def parse_uploaded_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses an uploaded document into structured text and metadata.
    Supports PDF, DOCX, XLSX, CSV, JSON, TXT, MD, HTML.
    """
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    text = ""
    sections: List[Dict[str, str]] = []

    try:
        if ext == "pdf":
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                p_text = page.extract_text() or ""
                if p_text.strip():
                    pages_text.append(f"--- Page {idx + 1} ---\n{p_text.strip()}")
                    sections.append({"title": f"Page {idx + 1}", "content": p_text.strip()})
            text = "\n\n".join(pages_text)

        elif ext in ("docx", "doc"):
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paras = []
            for p in doc.paragraphs:
                p_str = p.text.strip()
                if p_str:
                    paras.append(p_str)
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        paras.append(" | ".join(row_data))
            text = "\n\n".join(paras)
            sections.append({"title": filename, "content": text})

        elif ext in ("xlsx", "xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet_texts = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    s_block = f"### Sheet: {sheetname}\n" + "\n".join(rows_text)
                    sheet_texts.append(s_block)
                    sections.append({"title": sheetname, "content": "\n".join(rows_text)})
            text = "\n\n".join(sheet_texts)

        elif ext == "csv":
            content_str = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(content_str))
            rows = []
            for r in reader:
                clean_r = [val.strip() for val in r if val.strip()]
                if clean_r:
                    rows.append(" | ".join(clean_r))
            text = "\n".join(rows)
            sections.append({"title": filename, "content": text})

        elif ext == "json":
            content_str = file_bytes.decode("utf-8", errors="replace")
            parsed_json = json.loads(content_str)
            text = json.dumps(parsed_json, indent=2)
            sections.append({"title": filename, "content": text})

        elif ext in ("html", "htm"):
            from bs4 import BeautifulSoup
            content_str = file_bytes.decode("utf-8", errors="replace")
            soup = BeautifulSoup(content_str, "html.parser")
            for s in soup(["script", "style", "nav", "footer"]):
                s.decompose()
            text = soup.get_text(separator="\n", strip=True)
            sections.append({"title": filename, "content": text})

        elif ext in ("txt", "md", "markdown"):
            text = file_bytes.decode("utf-8", errors="replace")
            sections.append({"title": filename, "content": text})

        else:
            raise ValueError(f"Unsupported file format: .{ext}. Supported formats: PDF, DOCX, XLSX, CSV, JSON, TXT, MD, HTML.")

    except Exception as e:
        raise ValueError(f"Failed to parse {filename}: {str(e)}")

    if not text.strip():
        raise ValueError(f"File {filename} contains no extractable text.")

    return {
        "filename": filename,
        "extension": ext,
        "size_bytes": len(file_bytes),
        "clean_text": text.strip(),
        "sections": sections
    }
